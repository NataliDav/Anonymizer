import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook


class TableFormatError(Exception):
    pass


def _cell_to_str(v):
    if v is None:
        return ""
    return str(v)


def _as_native(s):
    if not isinstance(s, str):
        return s
    t = s.strip()
    if not t:
        return s
    try:
        return int(t)
    except ValueError:
        pass
    try:
        f = float(t)
        return int(f) if f.is_integer() else f
    except (ValueError, OverflowError):
        return s


def read_xlsx(path):
    try:
        wb = load_workbook(str(path), data_only=True)
    except Exception as e:
        raise TableFormatError(f"Не удалось прочитать XLSX-файл: {e}") from e
    ws = wb.active
    records = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    while records and not any(c is not None and str(c).strip() for c in records[0]):
        records.pop(0)
    while records and not any(c is not None and str(c).strip() for c in records[-1]):
        records.pop()
    if not records:
        raise TableFormatError("Файл пуст")
    headers = [_cell_to_str(c).strip() for c in records[0]]
    width = len(headers)
    data = []
    for r in records[1:]:
        cells = [_cell_to_str(c) for c in r]
        data.append(cells[:width] + [""] * max(0, width - len(cells)))
    return headers, data


def write_xlsx(path, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append([_cell_to_str(h) for h in headers])
    for r in rows:
        ws.append([_as_native(c) for c in r])
    wb.save(str(path))


def _detect_delimiter(header_line):
    counts = {d: header_line.count(d) for d in (",", ";", "\t")}
    best = max(counts, key=counts.get)
    return best if counts[best] > 0 else ","


def read_csv(path):
    raw = Path(path).read_bytes()
    text = None
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise TableFormatError("Не удалось определить кодировку CSV (нужна UTF-8 или cp1251)")
    lines = [ln for ln in text.splitlines()]
    while lines and not lines[0].strip():
        lines.pop(0)
    if not lines:
        raise TableFormatError("Файл пуст")
    delim = _detect_delimiter(lines[0])
    records = [r for r in csv.reader(lines, delimiter=delim) if any(x.strip() for x in r)]
    if not records:
        raise TableFormatError("Файл пуст")
    headers = [h.strip() for h in records[0]]
    width = len(headers)
    data = []
    for r in records[1:]:
        data.append(r[:width] + [""] * max(0, width - len(r)))
    return headers, data


def write_csv(path, headers, rows):
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(headers)
        writer.writerows(rows)


def read_table(path):
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".xlsx":
        return read_xlsx(p)
    if suffix == ".csv":
        return read_csv(p)
    raise TableFormatError(f"Неподдерживаемый формат файла «{suffix}» (ожидается .xlsx или .csv)")


def write_table(path, headers, rows):
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".xlsx":
        write_xlsx(p, headers, rows)
    elif suffix == ".csv":
        write_csv(p, headers, rows)
    else:
        raise TableFormatError(f"Неподдерживаемый формат сохранения «{suffix}»")
